"""
parse-ig-sales-report
=====================
Azure Function (v1 programming model, HTTP trigger).

Called by a Power Automate flow that monitors igreports@aubreys.group
for emails whose subject matches the pattern NN_Daily (where NN is a
zero-padded location ID, e.g. "05_Daily").  Power Automate extracts the
location ID from the subject and passes both the ID and the base64-
encoded Excel attachment in the JSON body.

Request body (JSON)
-------------------
{
    "location_id":  5,
    "report_data":  "<base64-encoded .xlsx bytes>"
}

Behaviour
---------
1. Decodes and opens the Excel workbook in memory.
2. Parses the "Processed Business Period" header line to get the report
   start/end datetimes.
3. Derives a business date per row using a 4:00 AM cutoff:
       closedDateTime >= 04:00  ->  businessDate = closedDateTime.date()
       closedDateTime <  04:00  ->  businessDate = closedDateTime.date() - 1 day
   This correctly handles both single-day and multi-day (back-fill) reports.
4. Checks both ig.hs_sales_landing AND ig.hs_sales for any existing rows
   covering the same locId + business date range.  Returns HTTP 409 if a
   conflict is found so Power Automate can alert and stop.
5. Extracts all data rows (skipping header, section-header, and footer
   rows), converts Excel serial date numbers to datetimes, and uploads a
   JSON array to Azure Blob Storage -- one blob per business date.
6. Returns a JSON summary on success.

The downstream ADF/Power Automate pipeline reads the blob and bulk-inserts
into ig.hs_sales_landing.

Environment variables required
-------------------------------
SQL_SERVER              Azure SQL server hostname
SQL_DATABASE            Database name
SQL_USER                SQL login username
SQL_PASSWORD            SQL login password
DATALAKE_IG_SALES_DATA  Full base URL for the sales blob folder,
                        e.g. https://aubdlakegen2.blob.core.windows.net/cherokee/ig_sales/data
DATALAKE_SAS            Shared access signature (existing, used across all datalake functions)
"""

import azure.functions as func
import base64
import io
import json
import logging
import os
import re
import pymssql
from datetime import datetime, timedelta, date
from decimal import Decimal, InvalidOperation
from openpyxl import load_workbook
from azure.storage.blob import BlobClient, ContentSettings


# ── Constants ──────────────────────────────────────────────────────────────────

EXCEL_EPOCH = datetime(1899, 12, 30)
BUSINESS_DAY_START_HOUR = 4

FOOTER_MARKERS = ("Profit Center Total", "Grand Total", "# of transactions:")


# ── Helpers ────────────────────────────────────────────────────────────────────

def excel_serial_to_datetime(serial):
    """Convert an Excel date serial number to a Python datetime."""
    if serial is None:
        return None
    try:
        return EXCEL_EPOCH + timedelta(days=float(serial))
    except (TypeError, ValueError):
        return None


def business_date_from_closed(closed_dt, fallback):
    """
    Return the business date for a transaction.
    Transactions before 4:00 AM belong to the previous business day.
    """
    if closed_dt is None:
        return fallback
    if closed_dt.hour < BUSINESS_DAY_START_HOUR:
        return (closed_dt - timedelta(days=1)).date()
    return closed_dt.date()


def safe_decimal(value):
    if value is None or str(value).strip() == "":
        return None
    try:
        return float(Decimal(str(value)))
    except (InvalidOperation, ValueError):
        return None


def safe_int(value):
    if value is None or str(value).strip() == "":
        return None
    try:
        return int(value)
    except (TypeError, ValueError):
        return None


def is_footer_row(row):
    first = str(row[0] or "").strip()
    return any(first.startswith(m) for m in FOOTER_MARKERS)


def is_section_header_row(row):
    """Rows like '    Profit Center:Aubrey's Maryville(959)' between data blocks."""
    first = str(row[0] or "").strip()
    return first.startswith("Profit Center:")


# ── Report parsing ─────────────────────────────────────────────────────────────

def parse_report_period(ws_rows):
    """
    Extract start/end datetimes from the report header.

    Expected header row (index 2):
        "Processed Business Period Starting 5/27/2026 4:00 AM and Ending 5/28/2026 3:59 AM"
    """
    header_text = ""
    for cell in ws_rows[2]:
        if cell is not None:
            header_text = str(cell)
            break

    pattern = (
        r"Starting\s+(\d{1,2}/\d{1,2}/\d{4}\s+\d{1,2}:\d{2}\s+[AP]M)"
        r"\s+and\s+Ending\s+(\d{1,2}/\d{1,2}/\d{4}\s+\d{1,2}:\d{2}\s+[AP]M)"
    )
    match = re.search(pattern, header_text, re.IGNORECASE)
    if not match:
        raise ValueError(
            f"Could not parse business period from report header. "
            f"Header text found: {header_text!r}"
        )

    fmt = "%m/%d/%Y %I:%M %p"
    start_dt = datetime.strptime(match.group(1).strip(), fmt)
    end_dt   = datetime.strptime(match.group(2).strip(), fmt)
    return start_dt, end_dt


def find_column_header_row(ws_rows):
    """Return the index of the row containing 'Closed Date/Time'."""
    for i, row in enumerate(ws_rows):
        if row[0] is not None and "Closed Date" in str(row[0]):
            return i
    raise ValueError("Could not find column header row ('Closed Date/Time') in report.")


def parse_data_rows(ws_rows, loc_id, report_start_dt, report_end_dt):
    """
    Parse all data rows from the workbook.

    Returns a dict keyed by business date so the caller can write one
    blob per business date -- important for multi-day back-fill reports.
    """
    fallback_date    = report_start_dt.date()
    upload_ts        = datetime.utcnow().isoformat()
    header_row_idx   = find_column_header_row(ws_rows)
    rows_by_date     = {}

    for row in ws_rows[header_row_idx + 1:]:
        # Skip entirely empty rows
        if all(cell is None or str(cell).strip() == "" for cell in row):
            continue
        # Footer reached -- stop
        if is_footer_row(row):
            break
        # Section-header rows appear in multi-profit-centre reports
        if is_section_header_row(row):
            continue

        # Column order: closedDateTime, checkNo, ctId, mpId, serverId, cashierId,
        #               grossRevenue, discount, tax, gratSvcChg, tip, roundedAmount,
        #               checkTotal, tenderId, tender, changeAmt, breakage, netTender
        closed_dt = excel_serial_to_datetime(row[0])
        biz_date  = business_date_from_closed(closed_dt, fallback_date)

        record = {
            "locId":          loc_id,
            "businessDate":   biz_date.isoformat(),
            "reportStartDt":  report_start_dt.isoformat(),
            "reportEndDt":    report_end_dt.isoformat(),
            "closedDateTime": closed_dt.isoformat() if closed_dt else None,
            "checkNo":        safe_int(row[1]),
            "ctId":           safe_int(row[2]),
            "mpId":           safe_int(row[3]),
            "serverId":       safe_int(row[4]),
            "cashierId":      safe_int(row[5]),
            "grossRevenue":   safe_decimal(row[6]),
            "discount":       safe_decimal(row[7]),
            "tax":            safe_decimal(row[8]),
            "gratSvcChg":     safe_decimal(row[9]),
            "tip":            safe_decimal(row[10]),
            "roundedAmount":  safe_decimal(row[11]),
            "checkTotal":     safe_decimal(row[12]),
            "tenderId":       safe_int(row[13]),
            "tender":         safe_decimal(row[14]),
            "changeAmt":      safe_decimal(row[15]),
            "breakage":       safe_decimal(row[16]),
            "netTender":      safe_decimal(row[17]),
            "tpl_upload":     upload_ts,
        }

        rows_by_date.setdefault(biz_date, []).append(record)

    return rows_by_date


# ── Database ───────────────────────────────────────────────────────────────────

def get_db_connection():
    return pymssql.connect(
        server=os.environ["SQL_SERVER"],
        database=os.environ["SQL_DATABASE"],
        user=os.environ["SQL_USER"],
        password=os.environ["SQL_PASSWORD"],
    )


def check_existing_data(loc_id, business_dates):
    """
    Return the subset of business_dates that already have data in either
    ig.hs_sales_landing or ig.hs_sales.
    """
    if not business_dates:
        return []

    placeholders = ", ".join(["%s"] * len(business_dates))
    date_strs    = [d.isoformat() for d in business_dates]

    query = f"""
        SELECT DISTINCT businessDate
        FROM (
            SELECT businessDate FROM ig.hs_sales_landing
            WHERE locId = %s AND businessDate IN ({placeholders})
            UNION ALL
            SELECT businessDate FROM ig.hs_sales
            WHERE locId = %s AND businessDate IN ({placeholders})
        ) combined
    """
    params = [loc_id] + date_strs + [loc_id] + date_strs

    with get_db_connection() as conn:
        cursor = conn.cursor()
        cursor.execute(query, params)
        rows = cursor.fetchall()

    return [row[0] for row in rows]


# ── Blob upload ────────────────────────────────────────────────────────────────

def upload_to_blob(loc_id, biz_date, records):
    """Upload a JSON array to blob storage using SAS and return the blob path."""
    base_url  = os.environ["DATALAKE_IG_SALES_DATA"].rstrip("/")
    sas_token = os.environ["DATALAKE_SAS"].lstrip("?")
    timestamp = datetime.utcnow().strftime("%Y%m%dT%H%M%SZ")
    blob_name = (
        f"loc{loc_id:02d}/{biz_date.isoformat()}/"
        f"hs_sales_{loc_id:02d}_{biz_date.isoformat()}_{timestamp}.json"
    )

    blob_client = BlobClient.from_blob_url(f"{base_url}/{blob_name}?{sas_token}")
    blob_client.upload_blob(
        json.dumps(records, default=str),
        overwrite=True,
        content_settings=ContentSettings(content_type="application/json"),
    )

    # Return path without SAS for logging/response
    return f"{base_url}/{blob_name}"


# ── Entry point ────────────────────────────────────────────────────────────────

def main(req: func.HttpRequest) -> func.HttpResponse:
    logging.info("parse-ig-sales-report: triggered")

    # Parse request body
    try:
        body = req.get_json()
    except ValueError:
        return func.HttpResponse("Request body must be valid JSON.", status_code=400)

    location_id     = body.get("location_id")
    report_data_b64 = body.get("report_data")

    if location_id is None or not report_data_b64:
        return func.HttpResponse(
            "Request body must include 'location_id' (int) and 'report_data' (base64 string).",
            status_code=400,
        )

    try:
        location_id = int(location_id)
        excel_bytes = base64.b64decode(report_data_b64)
    except Exception as e:
        return func.HttpResponse(f"Invalid input: {e}", status_code=400)

    # Open workbook
    try:
        wb      = load_workbook(io.BytesIO(excel_bytes), read_only=True, data_only=True)
        ws      = wb.active
        ws_rows = list(ws.iter_rows(values_only=True))
        wb.close()
    except Exception as e:
        logging.error(f"parse-ig-sales-report: workbook parse failed: {e}")
        return func.HttpResponse(f"Failed to open Excel file: {e}", status_code=422)

    # Parse report period from header
    try:
        report_start_dt, report_end_dt = parse_report_period(ws_rows)
    except ValueError as e:
        logging.error(f"parse-ig-sales-report: header parse failed: {e}")
        return func.HttpResponse(str(e), status_code=422)

    logging.info(
        f"parse-ig-sales-report: locId={location_id}  "
        f"period={report_start_dt} -> {report_end_dt}"
    )

    # Parse data rows
    try:
        rows_by_date = parse_data_rows(ws_rows, location_id, report_start_dt, report_end_dt)
    except Exception as e:
        logging.error(f"parse-ig-sales-report: data parse failed: {e}")
        return func.HttpResponse(f"Failed to parse report rows: {e}", status_code=422)

    if not rows_by_date:
        return func.HttpResponse(
            json.dumps({"status": "skipped", "reason": "No data rows found in report."}),
            mimetype="application/json",
            status_code=200,
        )

    all_business_dates = sorted(rows_by_date.keys())

    # Duplicate check against both landing and archive tables
    try:
        conflicting_dates = check_existing_data(location_id, all_business_dates)
    except Exception as e:
        logging.error(f"parse-ig-sales-report: DB duplicate check failed: {e}")
        return func.HttpResponse(f"Database check failed: {e}", status_code=500)

    if conflicting_dates:
        conflict_strs = [
            d.isoformat() if isinstance(d, date) else str(d)
            for d in conflicting_dates
        ]
        msg = (
            f"Data already exists for locId={location_id} on: "
            f"{', '.join(conflict_strs)}. Upload skipped."
        )
        logging.warning(f"parse-ig-sales-report: {msg}")
        return func.HttpResponse(
            json.dumps({"status": "conflict", "reason": msg, "conflictDates": conflict_strs}),
            mimetype="application/json",
            status_code=409,
        )

    # Upload one blob per business date
    results = []
    for biz_date, records in rows_by_date.items():
        try:
            blob_path = upload_to_blob(location_id, biz_date, records)
            logging.info(
                f"parse-ig-sales-report: uploaded {len(records)} rows -> {blob_path}"
            )
            results.append({
                "businessDate": biz_date.isoformat(),
                "rowCount":     len(records),
                "blobPath":     blob_path,
            })
        except Exception as e:
            logging.error(f"parse-ig-sales-report: blob upload failed for {biz_date}: {e}")
            return func.HttpResponse(
                f"Blob upload failed for {biz_date}: {e}", status_code=500
            )

    return func.HttpResponse(
        json.dumps({
            "status":      "success",
            "locId":       location_id,
            "reportStart": report_start_dt.isoformat(),
            "reportEnd":   report_end_dt.isoformat(),
            "uploads":     results,
        }),
        mimetype="application/json",
        status_code=200,
    )